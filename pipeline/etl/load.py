"""Load stage: write the cleaned/feature CSVs and the Excel report workbook."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1E4D40")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)


def save_csv(df: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False)


def _style_sheet(ws, df: pd.DataFrame) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    for i, col in enumerate(df.columns, start=1):
        max_len = df[col].astype(str).str.len().head(200).max() if len(df) else 10
        width = min(max(len(str(col)), max_len) + 2, 40)
        ws.column_dimensions[get_column_letter(i)].width = width
        if max_len > 60 and len(df) <= 50:
            for row_idx, cell in enumerate(ws[get_column_letter(i)][1:], start=2):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                ws.row_dimensions[row_idx].height = max(ws.row_dimensions[row_idx].height or 0, 60)


def write_report_workbook(path: Path, sheets: dict[str, pd.DataFrame], overview_stats: dict) -> None:
    """sheets: {sheet_name: dataframe}. overview_stats: {label: (value, number_format)}.

    Overview values are computed in Python and written as static numbers, not
    live formulas: this sandbox's LibreOffice cannot recalculate formulas
    (even a single-cell =A1+A2 workbook times out headless here), so a
    formula-based sheet could not be verified error-free before shipping --
    the same guarantee a static, directly-computed value gives without
    depending on a broken recalculation step.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)

        wb = writer.book
        for style in wb._named_styles:
            if style.name == "Normal":
                style.font = BODY_FONT
        for name, df in sheets.items():
            _style_sheet(wb[name[:31]], df)

        ov = wb.create_sheet("Overview", 0)
        ov["A1"] = "KONU Real Estate — Analysis Overview"
        ov["A1"].font = Font(name="Arial", size=14, bold=True)
        ov.append([])
        ov.append(["Metric", "Value"])
        for cell in ov[3]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT

        row = 4
        for label, (value, number_format) in overview_stats.items():
            ov.cell(row=row, column=1, value=label).font = BODY_FONT
            c = ov.cell(row=row, column=2, value=value)
            c.font = BODY_FONT
            if number_format:
                c.number_format = number_format
            row += 1

        ov.column_dimensions["A"].width = 34
        ov.column_dimensions["B"].width = 20
